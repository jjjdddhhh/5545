# 매뉴얼, 수치 대조(C08), 일정 배치(C09), CSV·ICS·DOCX 내보내기 테스트(DB 없이).
import io
from datetime import date
from types import SimpleNamespace

from app.export import csv_ics, docx_export
from app.llm import llm_client
from app.llm.prompt_store import file_prompt
from app.pipeline import checks, manual, schedule
from app.pipeline.llm_step import ListRecorder
from tests.fake_llm import FakeOllama, default_answer

SETTING = SimpleNamespace(audience="신입 사원", difficulty="beginner", output_language="ko", tone=None, keywords=[])
PARAS = [{"id": "p1", "kind": "body", "text": "처음에는 낮은 속도로 3초간 자리를 잡는다. 배터리는 40℃ 이하에서 보관한다."},
         {"id": "p2", "kind": "body", "text": "배터리는 매주 한 번 점검한다. 1,000회 충전 뒤에는 교체한다."}]


def test_extract_numbers_units_and_ordinals():
    assert checks.extract_numbers("3초간 잡고 40℃ 이하, 2단계에서 1,000회") == [("3", "초"), ("40", "°C"), ("1000", "회")]
    assert checks.extract_numbers("3번째 장면") == []                       # 순서 숫자는 수치가 아니다


def test_unsupported_numbers():
    ev = PARAS[0]["text"]
    assert checks.unsupported_numbers("낮은 속도로 3초 동안 잡는다", ev) == []
    assert checks.unsupported_numbers("5초 동안 잡고 40°C 이하에 둔다", ev) == ["5초"]  # ℃와 °C는 같게 본다
    assert checks.unsupported_numbers("3분 동안", ev) == ["3분"]             # 값이 같아도 단위가 다르면 실패


def test_layout_and_check_schedule():
    items = schedule.layout([{"seq": 2, "title": "보관", "duration_days": 0, "interval_days": 7},
                             {"seq": 1, "title": "점검", "duration_days": 2, "interval_days": None}])
    assert [(i["title"], i["start_offset_day"], i["duration_days"]) for i in items] == [("점검", 0, 2), ("보관", 2, 1)]
    assert items[1]["interval_days"] == 7
    assert checks.check_schedule([{**i, "step_seq": i["step_seq"]} for i in items]).result == "pass"
    broken = [dict(items[0]), {**items[1], "start_offset_day": 0}]
    assert checks.check_schedule(broken).result == "fail"
    assert schedule.clamp_days(9999) == 365 and schedule.clamp_days(-3) == 1


def test_manual_retries_on_invented_number(monkeypatch):
    def invented_then_ok(kw, n):
        ans = default_answer("ManualOut", kw)
        if n == 1:
            ans["steps"][0]["instruction"] = "배터리를 10분 동안 충전한다."          # 원고에 없는 수치
        return ans
    fake = FakeOllama({"ManualOut": invented_then_ok})
    monkeypatch.setattr(llm_client.client, "chat", fake)
    res = manual.generate_manual(ListRecorder(), file_prompt("manual"), SETTING, "<source>구성안</source>", PARAS)
    assert res.retried and not res.failures
    assert "10분" in fake.calls[1]["messages"][1]["content"]


def test_manual_failure_kept_and_normalized(monkeypatch):
    def bad(kw, n):
        ans = default_answer("ManualOut", kw)
        ans["steps"][0]["instruction"] = "10분 동안 충전한다."
        ans["steps"][1]["duration_days"] = 0
        ans["steps"][1]["source_paragraphs"] = ["p1", "p77"]
        return ans
    monkeypatch.setattr(llm_client.client, "chat", FakeOllama({"ManualOut": bad}))
    res = manual.generate_manual(ListRecorder(), file_prompt("manual"), SETTING, "<source>x</source>", PARAS)
    codes = sorted(f.code for f in res.failures)
    assert codes == ["C03", "C08"]
    assert res.manual.steps[1].duration_days == 1 and res.manual.steps[1].source_paragraphs == ["p1"]


MANUAL = {"manual": {"id": 3, "run_id": 1, "audience": "신입 사원", "difficulty": "beginner", "title": "드릴 매뉴얼",
                     "intro": "소개", "created_at": "2026-10-01T00:00:00"},
          "steps": [{"id": 10, "seq": 1, "title": "점검", "instruction": "배터리를 끼운다.", "tip": "딸깍 소리",
                     "source_paragraphs": ["p1"], "edited_fields": [], "check_status": "pass"},
                    {"id": 11, "seq": 2, "title": "보관", "instruction": "케이스에 넣는다.", "tip": None,
                     "source_paragraphs": ["p2"], "edited_fields": [], "check_status": "pass"}],
          "cautions": [{"id": 1, "severity": "danger", "body": "헐렁한 장갑 금지", "source": "ai"}],
          "schedule": [{"id": 20, "seq": 1, "manual_step_id": 10, "title": "점검", "start_offset_day": 0,
                        "duration_days": 2, "interval_days": None, "note": None, "source": "ai"},
                       {"id": 21, "seq": 2, "manual_step_id": 11, "title": "보관", "start_offset_day": 2,
                        "duration_days": 1, "interval_days": 7, "note": "주간 점검", "source": "ai"}],
          "total_days": 3}


def test_csv_dates_from_start_date():
    text = csv_ics.build_csv(MANUAL, date(2026, 10, 5)).decode("utf-8-sig")
    lines = text.strip().split("\r\n")
    assert lines[0].startswith("순서,일정,시작(며칠째)")
    assert lines[1] == "1,점검,1,2,,2026-10-05,2026-10-06,1. 점검,"
    assert lines[2] == "2,보관,3,1,7,2026-10-07,2026-10-07,2. 보관,주간 점검"


def test_ics_events_and_repeat():
    from icalendar import Calendar
    cal = Calendar.from_ical(csv_ics.build_ics(MANUAL, date(2026, 10, 5), project_id=1))
    events = [c for c in cal.walk() if c.name == "VEVENT"]
    assert len(events) == 2
    assert events[0].decoded("dtstart") == date(2026, 10, 5) and events[0].decoded("dtend") == date(2026, 10, 7)
    assert events[1]["rrule"]["INTERVAL"] == [7] and events[1]["rrule"]["COUNT"] == [12]
    assert "도움말: 딸깍 소리" in str(events[0]["description"])


def test_docx_contains_storyboard_and_manual():
    import docx
    outline = {"outline": {"id": 1, "run_id": 1, "title": "드릴", "summary": "요약", "learning_objectives": ["목표"],
                           "created_at": "2026-10-01T00:00:00"},
               "scenes": [{"id": 1, "seq": 1, "title": "도입", "key_point": "k", "source_paragraphs": ["p1"],
                           "screen_description": "손 클로즈업", "visual_suggestion": "도식", "on_screen_text": "배터리",
                           "duration_sec": 30, "char_budget": 150, "edited_fields": [], "start_sec": 0,
                           "narration": {"id": 1, "body": "드릴을 켭니다.", "char_count": 6, "est_duration_sec": 1.2,
                                         "is_edited": False}, "cues": [], "check_status": "pass"}],
               "total_sec": 30}
    data = docx_export.build_docx("교육", outline, MANUAL)
    d = docx.Document(io.BytesIO(data))
    texts = "\n".join(p.text for p in d.paragraphs)
    assert "스토리보드 · 드릴" in texts and "매뉴얼 · 드릴 매뉴얼" in texts and "[위험] 헐렁한 장갑 금지" in texts
    assert d.tables[0].rows[1].cells[5].text == "드릴을 켭니다."
    assert d.tables[1].rows[2].cells[3].text == "1일, 7일마다 반복"
