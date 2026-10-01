# 매뉴얼·일정 API와 내보내기 통합 테스트(TEST_DATABASE_URL이 있을 때만).
import io
import json

from sqlalchemy import select

from app.db import models as m
from tests.test_runner_api import fake, setup_project, wait_run  # noqa: F401


def run_both(client):
    pid = setup_project(client, content_type="both")
    run = wait_run(client, client.post(f"/api/projects/{pid}/runs").json()["run_id"])
    assert run["status"] == "done", run["error_message"]
    return pid


def test_manual_generated_with_schedule(client, fake):
    pid = run_both(client)
    mv = client.get(f"/api/projects/{pid}/manual").json()
    assert [s["seq"] for s in mv["steps"]] == [1, 2]
    assert [(i["start_offset_day"], i["duration_days"], i["interval_days"]) for i in mv["schedule"]] == \
        [(0, 1, None), (1, 2, 7)]
    assert mv["schedule"][0]["manual_step_id"] == mv["steps"][0]["id"]
    assert mv["cautions"][0]["severity"] == "danger" and mv["total_days"] == 3


def test_patch_step_and_schedule_relayout(client, db, fake):
    pid = run_both(client)
    mv = client.get(f"/api/projects/{pid}/manual").json()
    step = mv["steps"][0]
    r = client.patch(f"/api/manual-steps/{step['id']}", json={"title": "배터리 체결 점검", "tip": ""})
    out = r.json()
    assert out["steps"][0]["title"] == "배터리 체결 점검" and out["steps"][0]["edited_fields"] == ["title"]
    assert out["schedule"][0]["title"] == "배터리 체결 점검"               # 일정 제목도 따라 바뀐다

    item = mv["schedule"][0]
    out = client.patch(f"/api/schedule-items/{item['id']}", json={"duration_days": 4}).json()
    assert [(i["start_offset_day"], i["duration_days"]) for i in out["schedule"]] == [(0, 4), (4, 2)]  # 뒤 항목이 밀린다
    assert out["schedule"][0]["source"] == "user"
    out = client.patch(f"/api/schedule-items/{mv['schedule'][1]['id']}", json={"interval_days": 0}).json()
    assert out["schedule"][1]["interval_days"] is None                    # 0은 반복 없음
    revs = db.scalars(select(m.Revision.entity_type)).all()
    assert revs.count("manual_step") == 1 and revs.count("schedule_item") == 2
    assert client.patch(f"/api/schedule-items/{item['id']}", json={"duration_days": 0}).status_code == 422


def test_exports_csv_ics_docx_json(client, fake):
    pid = run_both(client)
    r = client.get(f"/api/projects/{pid}/export", params={"format": "csv", "start_date": "2026-11-02"})
    assert r.status_code == 200 and "2026-11-02" in r.content.decode("utf-8-sig")
    r = client.get(f"/api/projects/{pid}/export", params={"format": "ics"})
    assert r.status_code == 200 and b"BEGIN:VCALENDAR" in r.content and b"RRULE:FREQ=DAILY" in r.content and b"INTERVAL=7" in r.content
    r = client.get(f"/api/projects/{pid}/export", params={"format": "docx"})
    import docx
    assert "스토리보드" in "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
    r = client.get(f"/api/projects/{pid}/export", params={"format": "docx", "kind": "manual"})
    assert "스토리보드" not in "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
    data = json.loads(client.get(f"/api/projects/{pid}/export", params={"format": "json"}).content)
    assert data["outline"]["scenes"] and data["manual"]["steps"] and data["setting"]["content_type"] == "both"


def test_video_only_has_no_manual(client, fake):
    pid = setup_project(client, content_type="video")
    wait_run(client, client.post(f"/api/projects/{pid}/runs").json()["run_id"])
    assert client.get(f"/api/projects/{pid}/manual").status_code == 404
    assert client.get(f"/api/projects/{pid}/export", params={"format": "ics"}).status_code == 404
